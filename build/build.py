"""Genera index.html a partir de la planilla (xlsx exportado) y data/fotos.json.
Uso: python3 build/build.py planilla.xlsx "27-09-2026"
     python3 build/build.py --rehacer "27-09-2026"   (reusa los datos del index.html actual; para cambios de diseño)
"""
import json, sys, html, os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
xlsx, fecha = sys.argv[1], sys.argv[2]
fotos = json.load(open(os.path.join(ROOT, 'data', 'fotos.json'), encoding='utf-8'))
def num(v):
    if v in (None, ''): return None
    try: return round(float(str(v).replace('$', '').replace('.', '').replace(',', '.')) if isinstance(v, str) else float(v))
    except ValueError: return None
items = []
if xlsx == '--rehacer':
    prev = open(os.path.join(ROOT, 'index.html'), encoding='utf-8').read()
    a = prev.index('id="catalogo">') + len('id="catalogo">')
    items = json.loads(prev[a:prev.index('</script>', a)].replace('<\\/', '</'))
    for it in items:
        it.pop('img', None)
        if it['n'] in fotos and os.path.exists(os.path.join(ROOT, fotos[it['n']])): it['img'] = fotos[it['n']]
else:
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True)['Productos']
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    for r in ws.iter_rows(min_row=2, values_only=True):
        nombre = r[col['Producto']]
        if not nombre: continue
        it = {'n': str(nombre).strip(), 'c': str(r[col['Categoría']]).strip()}
        p = num(r[col['Precio venta unidad ($)']]); k = num(r[col['Pack venta (unidades)']]); kp = num(r[col['Precio venta pack ($)']])
        if p: it['p'] = p
        if k and kp: it['k'] = int(k); it['kp'] = kp
        if it['n'] in fotos and os.path.exists(os.path.join(ROOT, fotos[it['n']])): it['img'] = fotos[it['n']]
        items.append(it)
con = sum(1 for i in items if 'p' in i)
note = '' if con else 'Aún no hay precios cargados. Los precios aparecerán aquí a medida que se completen.'
t = open(os.path.join(ROOT, 'build', 'template.html'), encoding='utf-8').read()
js = json.dumps(items, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
head = ('<!doctype html><html lang="es"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">'
        '<meta name="theme-color" content="#CD212A"><meta name="robots" content="noindex">'
        '<link rel="icon" href="data:image/svg+xml,%3Csvg%20xmlns=%22http://www.w3.org/2000/svg%22%20viewBox=%220%200%2064%2064%22%3E%3Ccircle%20cx=%2232%22%20cy=%2232%22%20r=%2232%22%20fill=%22%231C1A1A%22/%3E%3Ccircle%20fill=%22%23FFFFFF%22%20cx=%2232%22%20cy=%2232%22%20r=%2229%22/%3E%3Cpath%20d=%22M14.38%2046.78A23%2023%200%200%201%2012.93%2019.14%22%20fill=%22none%22%20stroke=%22%231C1A1A%22%20stroke-width=%223.6%22%20stroke-linecap=%22round%22/%3E%3Cpath%20d=%22M8.37%2016.06L16.56%2013.75L17.49%2022.22Z%22%20fill=%22%231C1A1A%22/%3E%3Cpath%20d=%22M49.62%2017.22A23%2023%200%200%201%2051.07%2044.86%22%20fill=%22none%22%20stroke=%22%231C1A1A%22%20stroke-width=%223.6%22%20stroke-linecap=%22round%22/%3E%3Cpath%20d=%22M55.63%2047.94L47.44%2050.25L46.51%2041.78Z%22%20fill=%22%231C1A1A%22/%3E%3Cpath%20fill=%22%23CD212A%22%20d=%22M29.6%209.5h4.8v2.2h-.6v7.6c0%201.6.9%202.7%202.4%204%201.6%201.4%202.6%203%202.6%205.4V50c0%201.7-1.3%203-3%203h-7.6c-1.7%200-3-1.3-3-3V28.7c0-2.4%201-4%202.6-5.4%201.5-1.3%202.4-2.4%202.4-4v-7.6h-.6Z%22/%3E%3Crect%20x=%2225.4%22%20y=%2235%22%20width=%2213.2%22%20height=%228%22%20fill=%22%23FFFFFF%22/%3E%3Crect%20x=%2225.4%22%20y=%2237.2%22%20width=%2213.2%22%20height=%223.6%22%20fill=%22%231C1A1A%22/%3E%3C/svg%3E">'
        '<link rel="apple-touch-icon" href="img/icono.png">'
        '<style>html{-webkit-text-size-adjust:100%}body{margin:0}[hidden]{display:none!important}img{max-width:100%}</style></head><body>')
out = head + t.replace('__DATA__', js).replace('__FECHA__', html.escape(fecha)).replace('__NOTE__', f'<p class="note">{html.escape(note)}</p>' if note else '') + '</body></html>'
open(os.path.join(ROOT, 'index.html'), 'w', encoding='utf-8').write(out)
print(f'{len(items)} productos, {con} con precio, {sum(1 for i in items if "img" in i)} con foto')
